from __future__ import annotations

import csv
import json
from dataclasses import asdict
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from .models import Listing, Product

DEFAULT_COLUMNS = [
    "Id",
    "AdStatus",
    "Category",
    "GoodsType",
    "AdType",
    "Condition",
    "Apparel",
    "ApparelType",
    "Brand",
    "Size",
    "Title",
    "Description",
    "Price",
    "Address",
    "ContactMethod",
    "ContactPhone",
    "ManagerName",
    "ImageUrls",
]

HEADER_ALIASES = {
    "Images": "ImageUrls",
    "Image URLs": "ImageUrls",
    "Фото": "ImageUrls",
    "Название": "Title",
    "Описание": "Description",
    "Цена": "Price",
    "Адрес": "Address",
    "Телефон": "ContactPhone",
    "Размер": "Size",
    "Бренд": "Brand",
}


def listing_row(listing: Listing, seller: dict, separator: str) -> dict[str, object]:
    return {
        "Id": listing.listing_id,
        "AdStatus": "Free",
        "Category": "Одежда, обувь, аксессуары",
        "GoodsType": "Мужская одежда",
        "AdType": seller["ad_type"],
        "Condition": seller["condition"],
        "Apparel": "Мужская одежда",
        "ApparelType": listing.apparel_type,
        "Brand": listing.product.brand,
        "Size": listing.size,
        "Title": listing.title,
        "Description": listing.description,
        "Price": listing.selling_price,
        "Address": seller["address"],
        "ContactMethod": seller["contact_method"],
        "ContactPhone": seller["contact_phone"],
        "ManagerName": seller["manager_name"],
        "ImageUrls": separator.join(listing.image_urls),
    }


def _find_header(ws) -> tuple[int, list[str]]:
    for row_number in range(1, min(ws.max_row, 30) + 1):
        values = [
            str(cell.value).strip() if cell.value is not None else "" for cell in ws[row_number]
        ]
        normalized = [HEADER_ALIASES.get(value, value) for value in values]
        if "Id" in normalized and ("Title" in normalized or "Description" in normalized):
            return row_number, normalized
    raise ValueError(f"Could not find Avito header row in sheet {ws.title!r}")


def export_xlsx(
    listings: list[Listing], path: Path, seller: dict, export: dict, template: Path | None = None
) -> None:
    rows = [listing_row(item, seller, export["image_separator"]) for item in listings]
    if template:
        workbook = load_workbook(template)
        candidates = [ws for ws in workbook.worksheets if "инструк" not in ws.title.lower()]
        if not candidates:
            raise ValueError("The template does not contain a data sheet")
        worksheet = candidates[0]
        header_row, headers = _find_header(worksheet)
        if worksheet.max_row > header_row:
            worksheet.delete_rows(header_row + 1, worksheet.max_row - header_row)
    else:
        workbook = Workbook()
        worksheet = workbook.active
        worksheet.title = "Объявления"
        header_row = 1
        headers = DEFAULT_COLUMNS
        worksheet.append(headers)
        worksheet.freeze_panes = "A2"
        worksheet.sheet_view.showGridLines = False
        fill = PatternFill("solid", fgColor="1F4E78")
        for cell in worksheet[1]:
            cell.fill = fill
            cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in rows:
        values = []
        for header in headers:
            canonical = HEADER_ALIASES.get(header, header)
            values.append(row.get(canonical, ""))
        worksheet.append(values)

    for row in worksheet.iter_rows(min_row=header_row + 1):
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=False)
    for index, header in enumerate(headers, start=1):
        canonical = HEADER_ALIASES.get(header, header)
        width = 16
        if canonical in {"Title", "ImageUrls"}:
            width = 45
        elif canonical == "Description":
            width = 70
        worksheet.column_dimensions[worksheet.cell(header_row, index).column_letter].width = width
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def export_xml(listings: list[Listing], path: Path, seller: dict) -> None:
    root = ET.Element("Ads", {"formatVersion": "3", "target": "Avito.ru"})
    for listing in listings:
        row = listing_row(listing, seller, "|")
        ad = ET.SubElement(root, "Ad")
        for tag in [
            "Id",
            "AdStatus",
            "Category",
            "GoodsType",
            "AdType",
            "Condition",
            "Apparel",
            "ApparelType",
            "Brand",
            "Size",
            "Title",
            "Description",
            "Price",
            "Address",
            "ContactMethod",
            "ContactPhone",
            "ManagerName",
        ]:
            value = row.get(tag)
            if value not in (None, ""):
                ET.SubElement(ad, tag).text = str(value)
        images = ET.SubElement(ad, "Images")
        for url in listing.image_urls:
            ET.SubElement(images, "Image", {"url": url})
    ET.indent(root, space="  ")
    path.parent.mkdir(parents=True, exist_ok=True)
    ET.ElementTree(root).write(path, encoding="utf-8", xml_declaration=True)


def export_products(products: list[Product], output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    payload = [asdict(product) for product in products]
    (output_dir / "products.json").write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    with (output_dir / "products.csv").open("w", encoding="utf-8-sig", newline="") as stream:
        writer = csv.DictWriter(
            stream,
            fieldnames=["sku", "name", "brand", "product_type", "source_price", "sizes", "url"],
        )
        writer.writeheader()
        for product in products:
            writer.writerow(
                {
                    "sku": product.sku,
                    "name": product.name,
                    "brand": product.brand,
                    "product_type": product.product_type,
                    "source_price": product.source_price,
                    "sizes": "|".join(product.sizes),
                    "url": f"https://magic-shopp.com/product/{product.sku}",
                }
            )


def write_validation_report(listings: list[Listing], seller: dict, path: Path) -> dict:
    errors: list[str] = []
    warnings: list[str] = []
    if not seller.get("address"):
        errors.append("Не заполнен seller.address")
    if not seller.get("contact_phone") and seller.get("contact_method") != "В сообщениях":
        errors.append("Не заполнен seller.contact_phone")
    ids = [listing.listing_id for listing in listings]
    if len(ids) != len(set(ids)):
        errors.append("Есть повторяющиеся Id объявлений")
    for listing in listings:
        if not listing.image_urls:
            errors.append(f"{listing.listing_id}: нет фото")
        if len(listing.title) > 100:
            errors.append(f"{listing.listing_id}: заголовок длиннее 100 символов")
        if len(listing.description) > 7500:
            errors.append(f"{listing.listing_id}: описание длиннее 7500 символов")
        if not listing.size:
            warnings.append(f"{listing.listing_id}: размер не заполнен")
    report = {
        "valid": not errors,
        "listings": len(listings),
        "errors": sorted(set(errors)),
        "warnings": sorted(set(warnings)),
    }
    path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return report
